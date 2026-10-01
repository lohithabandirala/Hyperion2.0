import io
import json
import logging
from typing import Dict, Any

logger = logging.getLogger("ntro.export")

class ExportProvider:
    @staticmethod
    def export_pptx(content_str: str, provenance_id: str = None) -> io.BytesIO:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        if provenance_id:
            prs.core_properties.keywords = provenance_id

        # Try to parse content as json slides
        slides_data = []
        try:
            parsed = json.loads(content_str)
            if isinstance(parsed, dict) and "slides" in parsed:
                slides_data = parsed["slides"]
            elif isinstance(parsed, list):
                slides_data = parsed
        except Exception:
            pass

        if not slides_data:
            # Fallback slide from plain text
            slides_data = [
                {
                    "slide_number": 1,
                    "title": "NTRO Strategic Content Transformation",
                    "subtitle": "Generated Briefing",
                    "content": [line for line in content_str.split("\n") if line.strip()][:5],
                    "speaker_notes": "Official NTRO AI Platform Output"
                }
            ]

        for s in slides_data:
            blank_layout = prs.slide_layouts[6] # blank slide
            slide = prs.slides.add_slide(blank_layout)
            
            # Header title box
            title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.33), Inches(1.2))
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = s.get("title", f"Slide {s.get('slide_number', 1)}")
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(15, 23, 42) # Slate-900

            if s.get("subtitle"):
                p2 = tf.add_paragraph()
                p2.text = s.get("subtitle")
                p2.font.size = Pt(16)
                p2.font.color.rgb = RGBColor(37, 99, 235) # Blue-600

            # Content box
            content_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(4.5))
            ctf = content_box.text_frame
            ctf.word_wrap = True
            
            bullets = s.get("content", [])
            for idx, b in enumerate(bullets):
                cp = ctf.add_paragraph() if idx > 0 else ctf.paragraphs[0]
                cp.text = f"• {b}"
                cp.font.size = Pt(18)
                cp.font.color.rgb = RGBColor(51, 65, 85) # Slate-700
                cp.space_after = Pt(14)
                
            # Speaker notes
            if s.get("speaker_notes"):
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.text = s.get("speaker_notes")

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def export_docx(title: str, content_str: str, provenance_id: str = None) -> io.BytesIO:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        if provenance_id:
            doc.core_properties.keywords = provenance_id
            doc.core_properties.comments = f"NTRO Cryptographic Provenance ID: {provenance_id}"
        
        # Title
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"NATIONAL TECHNICAL RESEARCH ORGANISATION\n{title.upper()}")
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

        doc.add_paragraph("―" * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER

        for line in content_str.split("\n"):
            line_s = line.strip()
            if not line_s:
                continue
            if line_s.startswith("# "):
                h = doc.add_heading(line_s[2:], level=1)
                h.paragraph_format.space_before = Pt(12)
            elif line_s.startswith("## "):
                h = doc.add_heading(line_s[3:], level=2)
                h.paragraph_format.space_before = Pt(10)
            elif line_s.startswith("### "):
                h = doc.add_heading(line_s[4:], level=3)
            elif line_s.startswith("- ") or line_s.startswith("* "):
                p = doc.add_paragraph(line_s[2:], style='List Bullet')
            else:
                p = doc.add_paragraph(line_s)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def export_pdf(title: str, content_str: str) -> io.BytesIO:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#0f172a'),
            spaceAfter=6,
            alignment=1 # Center
        )
        sub_style = ParagraphStyle(
            'SubStyle',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#64748b'),
            spaceAfter=14,
            alignment=1
        )
        h2_style = ParagraphStyle(
            'H2Style',
            parent=styles['Heading2'],
            fontSize=13,
            textColor=colors.HexColor('#1e40af'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#334155'),
            leading=14,
            spaceAfter=6
        )

        story = []
        story.append(Paragraph("NATIONAL TECHNICAL RESEARCH ORGANISATION", sub_style))
        story.append(Paragraph(title.upper(), title_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceAfter=14))

        for line in content_str.split("\n"):
            line_s = line.strip()
            if not line_s:
                story.append(Spacer(1, 4))
                continue
            
            clean_line = line_s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if line_s.startswith("# "):
                story.append(Paragraph(clean_line[2:], h2_style))
            elif line_s.startswith("## "):
                story.append(Paragraph(clean_line[3:], h2_style))
            elif line_s.startswith("### "):
                story.append(Paragraph(f"<b>{clean_line[4:]}</b>", body_style))
            elif line_s.startswith("- ") or line_s.startswith("* "):
                story.append(Paragraph(f"&bull; {clean_line[2:]}", body_style))
            else:
                story.append(Paragraph(clean_line, body_style))

        doc.build(story)
        buffer.seek(0)
        return buffer
